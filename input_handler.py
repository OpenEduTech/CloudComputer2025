# input_handler.py
from pptx import Presentation
import os
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


def extract_text_from_pptx(file_path):
    """
    从 .pptx 文件中提取每一页的纯文本内容。
    
    Args:
        file_path (str): PPTX 文件的路径
    
    Returns:
        list[dict]: 每页一个字典，包含 'slide_index' 和 'text'
                    例如: [{'slide_index': 0, 'text': '标题\n要点1\n要点2'}, ...]
                    
    Raises:
        FileNotFoundError: 文件不存在
        ValueError: 文件不是有效的 .pptx
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")
    
    if not file_path.lower().endswith('.pptx'):
        raise ValueError("仅支持 .pptx 格式的 PowerPoint 文件")

    try:
        prs = Presentation(file_path)
        slides_text = []

        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        slide_text.append(text)
            full_text = "\n".join(slide_text)
            if full_text:  # 忽略完全空白的幻灯片
                slides_text.append({
                    "slide_index": i,
                    "text": full_text
                })
        
        logger.info(f"成功从 {file_path} 提取 {len(slides_text)} 页非空内容")
        return slides_text

    except Exception as e:
        logger.error(f"解析 PPTX 文件失败: {e}")
        raise ValueError(f"无法解析 PPTX 文件: {e}")