"""
PPT文档解析器
"""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import fitz  # PyMuPDF
from PIL import Image
import pytesseract
import io
import os
import logging
import json
from typing import List, Dict, Any
from openai import AsyncOpenAI

from app.models.schemas import PPTSlide
from app.core.config import settings
from app.core.llm_validation import validate_outline_summary

logger = logging.getLogger(__name__)


class PPTParser:
    """PPT解析器"""
    
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.file_extension = os.path.splitext(file_path)[1].lower()
        self.client = AsyncOpenAI(
            api_key=settings.DEEPSEEK_API_KEY,
            base_url=settings.deepseek_api_base
        )
        self.model = settings.DEEPSEEK_MODEL
    
    async def parse(self) -> Dict[str, Any]:
        """解析PPT文件"""
        try:
            if self.file_extension in ['.ppt', '.pptx']:
                return await self._parse_pptx()
            elif self.file_extension == '.pdf':
                return await self._parse_pdf()
            else:
                raise ValueError(f"不支持的文件类型: {self.file_extension}")
        except Exception as e:
            logger.error(f"解析文件失败: {e}", exc_info=True)
            raise
    
    async def _parse_pptx(self) -> Dict[str, Any]:
        """解析PPTX文件"""
        logger.info(f"开始解析PPTX文件: {self.file_path}")
        
        prs = Presentation(self.file_path)
        slides_data = []
        
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_content = await self._extract_slide_content(slide, slide_idx)
            slides_data.append(slide_content)
        
        metadata = {
            "slide_width": prs.slide_width,
            "slide_height": prs.slide_height,
            "core_properties": {
                "title": prs.core_properties.title,
                "author": prs.core_properties.author,
                "created": str(prs.core_properties.created) if prs.core_properties.created else None,
            }
        }

        # 结构化目录（便于前端展示）
        metadata["outline"] = [
            {
                "slide_number": s.slide_number,
                "title": s.title,
                "subtitle": s.subtitle,
                "bullets": len(s.body or []),
                "images": len(s.images or []),
            }
            for s in slides_data
        ]

        if settings.OUTLINE_SUMMARY_ENABLED and metadata["outline"]:
            try:
                summary_raw, summary_list = await self._summarize_outline(metadata["outline"])
                metadata["outline_summary"] = summary_list
                if settings.LLM_VALIDATION_ENABLED:
                    metadata["llm_validation"] = validate_outline_summary(summary_raw, summary_list)
            except Exception as e:
                logger.warning(f"目录概括失败: {e}")
        
        logger.info(f"PPTX解析完成，共{len(slides_data)}页")
        
        return {
            "total_slides": len(slides_data),
            "slides": slides_data,
            "metadata": metadata
        }
    
    async def _extract_slide_content(self, slide, slide_number: int) -> PPTSlide:
        """提取单页PPT内容"""
        title = None
        subtitle = None
        content_list = []
        body_list = []
        images = []
        image_descriptions = []
        ocr_texts = []
        notes = None
        
        # 提取标题
        if slide.shapes.title:
            title = slide.shapes.title.text.strip()
        
        # 提取内容和图片
        text_blocks = self._collect_text_blocks(slide)

        for shape in slide.shapes:
            # 文本内容
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text and text != title:  # 避免重复标题
                    content_list.append(text)
            
            # 图片
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_info = f"图片_{slide_number}_{len(images)+1}"
                images.append(image_info)

                # OCR提取图片文字（可选）
                ocr_text = ""
                if settings.OCR_ENABLED:
                    try:
                        image_blob = shape.image.blob
                        ocr_text = self._extract_ocr_from_bytes(image_blob)
                        if ocr_text:
                            ocr_texts.append(ocr_text)
                    except Exception as e:
                        logger.warning(f"OCR提取失败: {e}")

                description = await self._build_image_description(image_info, ocr_text, title, subtitle)
                image_descriptions.append(description)
                content_list.append(f"【图片】{description}")
        
        # 提取备注
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            if notes_slide.notes_text_frame:
                notes = notes_slide.notes_text_frame.text.strip()
        
        # 解析层级结构：标题/副标题/正文
        if text_blocks:
            # 去掉标题文本
            filtered = [b for b in text_blocks if b["text"] != title]
            if filtered:
                # 以字号分层
                sizes = sorted({b["size"] for b in filtered}, reverse=True)
                if sizes:
                    subtitle_size = sizes[0]
                    subtitle_candidates = [b["text"] for b in filtered if b["size"] == subtitle_size]
                    subtitle = subtitle_candidates[0] if subtitle_candidates else None
                    body_list = [b["text"] for b in filtered if b["text"] != subtitle]
        if not body_list:
            body_list = [c for c in content_list if c and c != title and c != subtitle]

        return PPTSlide(
            slide_number=slide_number,
            title=title,
            subtitle=subtitle,
            content=content_list,
            body=body_list,
            images=images,
            image_descriptions=image_descriptions,
            ocr_texts=ocr_texts,
            notes=notes
        )
    
    async def _parse_pdf(self) -> Dict[str, Any]:
        """解析PDF文件（作为PPT的替代格式）"""
        logger.info(f"开始解析PDF文件: {self.file_path}")
        
        doc = fitz.open(self.file_path)
        slides_data = []
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            
            # 提取文本
            text = page.get_text()
            lines = [line.strip() for line in text.split('\n') if line.strip()]
            
            # 简单启发式：第一行作为标题，第二行作为副标题
            title = lines[0] if lines else None
            subtitle = lines[1] if len(lines) > 1 else None
            content = lines[2:] if len(lines) > 2 else (lines[1:] if len(lines) > 1 else [])
            
            # 提取图片信息
            image_list = page.get_images()
            images = [f"图片_{page_num+1}_{i+1}" for i in range(len(image_list))]

            # OCR提取（可选）：对整页渲染进行OCR
            ocr_texts = []
            image_descriptions = []
            if settings.OCR_ENABLED:
                try:
                    pix = page.get_pixmap(dpi=200)
                    img_data = pix.tobytes("png")
                    ocr_text = self._extract_ocr_from_bytes(img_data)
                    if ocr_text:
                        ocr_texts.append(ocr_text)
                        desc = await self._build_image_description(f"图片_{page_num+1}_1", ocr_text, title, subtitle)
                        image_descriptions.append(desc)
                except Exception as e:
                    logger.warning(f"PDF OCR提取失败: {e}")
            
            slides_data.append(PPTSlide(
                slide_number=page_num + 1,
                title=title,
                subtitle=subtitle,
                content=content + [f"【图片】{d}" for d in image_descriptions],
                body=content,
                images=images,
                image_descriptions=image_descriptions,
                ocr_texts=ocr_texts
            ))
        
        doc.close()
        
        logger.info(f"PDF解析完成，共{len(slides_data)}页")
        
        return {
            "total_slides": len(slides_data),
            "slides": slides_data,
            "metadata": {"source": "pdf"}
        }

    def _extract_ocr_from_bytes(self, image_bytes: bytes) -> str:
        """从图片字节流提取OCR文本"""
        if settings.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD

        with Image.open(io.BytesIO(image_bytes)) as img:
            text = pytesseract.image_to_string(img, lang="eng+chi_sim")
            return text.strip()

    async def _build_image_description(
        self,
        image_name: str,
        ocr_text: str,
        title: str | None,
        subtitle: str | None
    ) -> str:
        """构建图片描述（基于OCR文本 + LLM语义摘要）"""
        clean_text = (ocr_text or "").strip()
        if not clean_text:
            return f"{image_name}：未识别到文字内容"

        if settings.OCR_SEMANTIC_ENABLED:
            try:
                summary = await self._describe_image_semantics(clean_text, title, subtitle)
                if summary:
                    return f"{image_name}：{summary}"
            except Exception as e:
                logger.warning(f"图片语义摘要失败: {e}")

        preview = clean_text.replace("\n", " ")[:200]
        return f"{image_name}：{preview}"

    def _collect_text_blocks(self, slide) -> List[Dict[str, Any]]:
        """收集文本块与字号，便于分层"""
        blocks: List[Dict[str, Any]] = []
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame") or not shape.text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            max_size = 0.0
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text and run.font.size:
                        try:
                            max_size = max(max_size, float(run.font.size.pt))
                        except Exception:
                            continue
            blocks.append({"text": text, "size": max_size})
        return blocks

    async def _describe_image_semantics(self, ocr_text: str, title: str | None, subtitle: str | None) -> str:
        """使用LLM对OCR文本进行语义概括"""
        if not settings.DEEPSEEK_API_KEY:
            return ""
        title_text = title or "无"
        subtitle_text = subtitle or "无"
        prompt = f"""你是文档理解助手。请基于OCR文本给出图片语义解说，要求简短、准确。

页面标题：{title_text}
页面副标题：{subtitle_text}
OCR文本：{ocr_text}

输出要求：
- 20~60字
- 只输出一句话
"""

        response = await self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": "你是精炼的图像解说助手。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.3,
            max_tokens=120
        )
        return (response.choices[0].message.content or "").strip()

    async def _summarize_outline(self, outline: List[Dict[str, Any]]) -> tuple[str, List[str]]:
        """使用LLM对目录进行概括"""
        if not settings.DEEPSEEK_API_KEY:
            return "", []
        outline_text = "\n".join(
            f"第{o['slide_number']}页：{o.get('title') or '无标题'} {('- ' + o.get('subtitle')) if o.get('subtitle') else ''}"
            for o in outline
        )
        prompt = f"""你是课程助教。请将以下PPT目录概括为3-8条主题要点：

目录：
{outline_text}

输出JSON格式（仅输出JSON）：
{{"summary": ["主题1", "主题2", "主题3"]}}
"""
        response = await self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": "你是目录概括助手。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.2,
            max_tokens=200,
            response_format={"type": "json_object"}
        )
        raw = response.choices[0].message.content or ""
        try:
            start = raw.find("{")
            end = raw.rfind("}") + 1
            data = json.loads(raw[start:end]) if start != -1 and end > start else {}
            summary = data.get("summary") or []
            if isinstance(summary, list):
                summary = [str(s).strip() for s in summary if str(s).strip()]
            else:
                summary = []
            return raw, summary
        except Exception:
            return raw, []


class DocumentIndexer:
    """文档索引器 - 用于向量化存储"""
    
    def __init__(self, chroma_client):
        self.chroma_client = chroma_client
        self.collection_name = "ppt_slides"
    
    async def index_document(self, file_id: str, slides: List[PPTSlide]):
        """索引文档到向量数据库"""
        try:
            collection = self.chroma_client.get_or_create_collection(
                name=self.collection_name,
                metadata={"description": "PPT slides vector store"}
            )
            
            # 准备文档和元数据
            documents = []
            metadatas = []
            ids = []
            
            for slide in slides:
                # 合并标题和内容作为文档
                doc_text = f"{slide.title or ''}\n" + "\n".join(slide.content)
                if slide.image_descriptions:
                    doc_text += f"\n图片描述: {' | '.join(slide.image_descriptions)}"
                if slide.ocr_texts:
                    doc_text += f"\nOCR: {' '.join(slide.ocr_texts)}"
                if slide.notes:
                    doc_text += f"\n备注: {slide.notes}"
                
                documents.append(doc_text)
                metadatas.append({
                    "file_id": file_id,
                    "slide_number": slide.slide_number,
                    "title": slide.title or "",
                    "has_images": len(slide.images) > 0
                })
                ids.append(f"{file_id}_slide_{slide.slide_number}")
            
            # 添加到集合（可重复写入时使用upsert）
            collection.upsert(
                documents=documents,
                metadatas=metadatas,
                ids=ids
            )
            
            logger.info(f"文档{file_id}已索引，共{len(documents)}个切片")
            
        except Exception as e:
            logger.error(f"索引文档失败: {e}", exc_info=True)
            raise
    
    async def search_similar(self, query: str, file_id: str = None, top_k: int = 5):
        """搜索相似内容"""
        try:
            collection = self.chroma_client.get_collection(name=self.collection_name)
            
            where_filter = {"file_id": file_id} if file_id else None
            
            results = collection.query(
                query_texts=[query],
                n_results=top_k,
                where=where_filter
            )
            
            return results
            
        except Exception as e:
            logger.error(f"搜索失败: {e}", exc_info=True)
            raise
